"""
ReportFactory 测试（Sprint 4.8.4）
覆盖：学生名册 Excel、证书 docx/PDF、懒加载子服务、打印流程
"""

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent.parent / "src"))

from edu_system.services.report_factory import ReportFactory


@pytest.fixture
def factory():
    return ReportFactory()


@pytest.fixture
def excel_tpl(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "姓名：{{name}}"
    p = tmp_path / "roster.xlsx"
    wb.save(p)
    return p


@pytest.fixture
def docx_tpl(tmp_path):
    p = tmp_path / "cert.docx"
    from docx import Document

    doc = Document()
    doc.add_paragraph("兹证明 {{name}} 同学")
    doc.save(str(p))
    return p


class TestRoster:
    def test_gen_student_roster(self, factory, excel_tpl, tmp_path):
        data = {"name": "张三"}
        out = factory.gen_student_roster(excel_tpl, data, output_path=tmp_path / "out.xlsx")
        assert isinstance(out, bytes)
        assert len(out) > 0


class TestCertificate:
    def test_gen_certificate_docx(self, factory, docx_tpl, tmp_path):
        out = factory.gen_certificate(docx_tpl, {"name": "李四"}, tmp_path / "cert.docx")
        assert out.exists()

    def test_gen_certificate_batch(self, factory, docx_tpl, tmp_path):
        rows = [{"name": "A"}, {"name": "B"}]
        outs = factory.gen_and_print(docx_tpl, rows, tmp_path / "out", prefix="c")
        assert len(outs) == 2


class TestLazyLoading:
    def test_excel_lazy(self, factory):
        assert factory._excel is None
        _ = factory.excel
        assert factory._excel is not None

    def test_cert_lazy(self, factory):
        assert factory._cert is None
        _ = factory.cert
        assert factory._cert is not None

    def test_print_lazy(self, factory):
        assert factory._print is None
        _ = factory.printer
        assert factory._print is not None


class TestPrint:
    def test_print_files(self, factory, tmp_path):
        f1 = tmp_path / "a.pdf"
        f1.write_bytes(b"%PDF-1.4")
        # 无打印后端时返回 False
        results = factory.print_files([f1])
        assert str(f1) in results
